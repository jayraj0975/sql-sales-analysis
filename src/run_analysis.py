"""Run every query in sql/ against the Chinook database and build the reports.

Outputs
  reports/results/*.csv      one CSV per query
  reports/sales_analysis.xlsx  a workbook with a summary sheet, one sheet per query and charts
  reports/figures/*.png      three charts for the README
  reports/findings.md        findings, with every number computed from the query results

The SQL is the analysis; this file only runs it and formats the output.
"""

from __future__ import annotations

import sqlite3
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.chart import BarChart, LineChart, Reference  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
DB = ROOT / "data" / "chinook.sqlite"
SQL_DIR = ROOT / "sql"
REPORTS = ROOT / "reports"

BLUE, GREY = "#2a5bd7", "#8a93a3"


def connect(path: Path = DB) -> sqlite3.Connection:
    if not path.exists():
        raise FileNotFoundError(f"{path} not found. Run: python src/download_data.py")
    conn = sqlite3.connect(f"file:{path}?mode=ro", uri=True)  # read-only: the analysis never writes
    return conn


def query_files() -> list[Path]:
    return sorted(SQL_DIR.glob("*.sql"))


def query_name(path: Path) -> str:
    return path.stem.split("_", 1)[1]           # 01_yearly_revenue -> yearly_revenue


def query_title(path: Path) -> str:
    return path.read_text().splitlines()[0].lstrip("- ").strip()


def run_queries(conn: sqlite3.Connection) -> dict[str, pd.DataFrame]:
    return {query_name(p): pd.read_sql_query(p.read_text(), conn) for p in query_files()}


# ------------------------------------------------------------------ Excel
def write_excel(results: dict[str, pd.DataFrame], path: Path) -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = "Chinook sales analysis"
    summary["A1"].font = Font(size=14, bold=True)
    summary["A2"] = "Source: Chinook sample database (MIT licence). Every sheet is the output of one SQL query in sql/."
    summary["A4"], summary["B4"] = "Sheet", "Question"
    for c in ("A4", "B4"):
        summary[c].font = Font(bold=True)

    header_fill = PatternFill("solid", fgColor="2A5BD7")
    sheets = {}
    for row, p in enumerate(query_files(), start=5):
        name = query_name(p)
        summary.cell(row=row, column=1, value=name)
        summary.cell(row=row, column=2, value=query_title(p))
        df = results[name]
        ws = wb.create_sheet(name[:31])
        sheets[name] = ws
        ws.append(list(df.columns))
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for rec in df.itertuples(index=False):
            ws.append([None if pd.isna(v) else v for v in rec])
        ws.freeze_panes = "A2"
        for i, col in enumerate(df.columns, start=1):
            width = max(len(str(col)), *(len(str(v)) for v in df[col].head(60))) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(width, 40)
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 100

    def bar(ws, title, cat_col, val_col, rows, anchor, y_title):
        ch = BarChart()
        ch.title, ch.y_axis.title, ch.legend = title, y_title, None
        ch.add_data(Reference(ws, min_col=val_col, min_row=1, max_row=rows + 1), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=cat_col, min_row=2, max_row=rows + 1))
        ch.height, ch.width = 8, 16
        ch.y_axis.scaling.min = 0          # bars must start at zero or small differences look large
        ch.x_axis.delete = ch.y_axis.delete = False   # openpyxl hides axes in Excel unless told otherwise
        ws.add_chart(ch, anchor)

    y = results["yearly_revenue"]
    bar(sheets["yearly_revenue"], "Revenue by year", 1, 4, len(y), "I2", "Revenue ($)")

    m = results["monthly_revenue"]
    ws = sheets["monthly_revenue"]
    line = LineChart()
    line.title, line.y_axis.title = "Monthly revenue and 3-month average", "Revenue ($)"
    line.add_data(Reference(ws, min_col=3, min_row=1, max_row=len(m) + 1, max_col=4), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(m) + 1))
    line.height, line.width = 8, 22
    line.y_axis.scaling.min = 0
    line.x_axis.delete = line.y_axis.delete = False
    for series in line.series:
        series.smooth = False              # draw the data as it is, not a fitted curve
    ws.add_chart(line, "F2")

    bar(sheets["revenue_by_country"], "Revenue by country (top 10)", 1, 4, min(10, len(results["revenue_by_country"])), "I2", "Revenue ($)")
    bar(sheets["revenue_by_genre"], "Revenue by genre (top 10)", 1, 4, min(10, len(results["revenue_by_genre"])), "I2", "Revenue ($)")
    for ws in wb.worksheets:                   # print each sheet, and its chart, on one page width
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    wb.save(path)


# ---------------------------------------------------------------- figures
def style(ax) -> None:
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(alpha=0.25)


def write_figures(results: dict[str, pd.DataFrame], out: Path) -> None:
    out.mkdir(parents=True, exist_ok=True)

    m = results["monthly_revenue"]
    fig, ax = plt.subplots(figsize=(9, 3.8))
    ax.bar(m["month"], m["revenue"], color=BLUE, alpha=0.35, label="Monthly revenue")
    ax.plot(m["month"], m["moving_avg_3m"], color=BLUE, lw=2, label="3-month average")
    ax.set_xticks(range(0, len(m), 6), m["month"][::6], rotation=45, ha="right")
    ax.set(title="Monthly revenue", ylabel="Revenue ($)")
    ax.legend(frameon=False)
    style(ax)
    fig.tight_layout()
    fig.savefig(out / "monthly_revenue.png", dpi=130)
    plt.close(fig)

    c = results["revenue_by_country"].head(8).iloc[::-1]
    fig, ax = plt.subplots(figsize=(7, 4))
    ax.barh(c["country"], c["revenue"], color=BLUE)
    ax.set(title="Revenue by billing country (top 8)", xlabel="Revenue ($)")
    style(ax)
    fig.tight_layout()
    fig.savefig(out / "revenue_by_country.png", dpi=130)
    plt.close(fig)

    p = results["customer_pareto"]
    fig, ax = plt.subplots(figsize=(6.5, 4))
    x = 100 * p["customer_rank"] / len(p)
    ax.plot(x, p["cumulative_pct"], color=BLUE, lw=2)
    ax.plot([0, 100], [0, 100], color=GREY, ls="--", lw=1.2, label="Perfectly even")
    ax.set(title="Customer revenue concentration", xlabel="Share of customers, %", ylabel="Cumulative share of revenue, %")
    ax.legend(frameon=False)
    style(ax)
    fig.tight_layout()
    fig.savefig(out / "customer_pareto.png", dpi=130)
    plt.close(fig)


# --------------------------------------------------------------- findings
def md_table(df: pd.DataFrame) -> str:
    def fmt(v):
        if isinstance(v, float):
            return "" if pd.isna(v) else f"{v:,.2f}"
        return str(v)

    cells = [list(map(str, df.columns))] + [[fmt(v) for v in r] for r in df.itertuples(index=False)]
    lines = ["| " + " | ".join(cells[0]) + " |", "|" + "|".join("---" for _ in cells[0]) + "|"]
    lines += ["| " + " | ".join(r) + " |" for r in cells[1:]]
    return "\n".join(lines)


def key_numbers(r: dict[str, pd.DataFrame]) -> dict:
    y, p, c, g = r["yearly_revenue"], r["customer_pareto"], r["revenue_by_country"], r["revenue_by_genre"]
    cat, coh, seg = r["catalog_coverage"], r["cohort_activity"], r["rfm_segments"]
    total = float(y["revenue"].sum())
    mon = r["monthly_revenue"]["revenue"]
    return {
        "mode_month_value": float(mon.mode()[0]),
        "months_at_mode": int((mon == mon.mode()[0]).sum()),
        "n_months": len(mon),
        "total_revenue": total,
        "invoices": int(y["invoices"].sum()),
        "customers": len(p),
        "years": f"{y['year'].iloc[0]} to {y['year'].iloc[-1]}",
        "year_min": float(y["revenue"].min()), "year_max": float(y["revenue"].max()),
        "yoy": [v for v in y["yoy_pct"].tolist() if pd.notna(v)],
        "aov": total / int(y["invoices"].sum()),
        "n_for_50": int((p["cumulative_pct"] < 50).sum() + 1),
        "n_for_80": int((p["cumulative_pct"] < 80).sum() + 1),
        "top10_share": float(p["share_pct"].head(10).sum()),
        "top_country": c.iloc[0]["country"], "top_country_share": float(c.iloc[0]["share_pct"]),
        "top3_country_share": float(c["share_pct"].head(3).sum()),
        "rpc_min": float(c["revenue_per_customer"].min()), "rpc_max": float(c["revenue_per_customer"].max()),
        "n_countries": len(c), "n_genres": len(g),
        "rpc5_min": float(c["revenue_per_customer"].head(5).min()), "rpc5_max": float(c["revenue_per_customer"].head(5).max()),
        "freq_mode": int(seg["frequency"].mode()[0]),
        "freq_same": int((seg["frequency"] == seg["frequency"].mode()[0]).sum()),
        "one_customer_countries": int((c["customers"] == 1).sum()),
        "top_genre": g.iloc[0]["genre"], "top_genre_share": float(g.iloc[0]["share_pct"]),
        "top3_genre_share": float(g["share_pct"].head(3).sum()),
        "catalogue": int(cat["tracks_in_catalogue"].sum()),
        "never_sold": int((cat["tracks_in_catalogue"] - cat["tracks_ever_sold"]).sum()),
        "segments": seg["segment"].value_counts().to_dict(),
        "cohorts": coh,
    }


def write_findings(r: dict[str, pd.DataFrame], path: Path) -> None:
    k = key_numbers(r)
    never = 100 * k["never_sold"] / k["catalogue"]
    coh = k["cohorts"]
    first = coh["cohort_year"].min()
    later = coh[(coh["cohort_year"] == first) & (coh["activity_year"] > first)]
    text = f"""# Findings

Generated by `src/run_analysis.py`. Every number below is computed from the query results.

**About the data.** Chinook is a *sample* database for a fictional digital music store
({k['years']}, {k['customers']} customers). It is built for teaching SQL, and its sales are very
evenly spread, so several results below are "flat" or "even" rather than dramatic. The value
of this project is the SQL and its checks, not a discovery about a real business.

## Headline numbers

| Measure | Value |
|---|---|
| Revenue | ${k['total_revenue']:,.2f} over {k['invoices']} invoices |
| Customers with at least one invoice | {k['customers']} |
| Average order value | ${k['aov']:.2f} |
| Revenue per year | ${k['year_min']:,.2f} to ${k['year_max']:,.2f} |

## What the queries show

1. **Revenue is flat.** Yearly revenue stays between ${k['year_min']:,.0f} and ${k['year_max']:,.0f}; year-over-year
   changes run from {min(k['yoy']):+.1f}% to {max(k['yoy']):+.1f}%. There is no growth trend to explain. In {k['months_at_mode']} of
   {k['n_months']} months revenue is *exactly* ${k['mode_month_value']:,.2f}, which is the signature of generated data, not of a real store.
2. **No small group of customers dominates revenue.** It takes {k['n_for_50']} customers to reach 50% of revenue and
   {k['n_for_80']} of {k['customers']} to reach 80%; the top 10 customers hold {k['top10_share']:.1f}%.
3. **{k['top_country']} is the largest market** at {k['top_country_share']:.1f}% of revenue, and the top three countries together
   are {k['top3_country_share']:.1f}%. Among the five largest countries revenue per customer is ${k['rpc5_min']:.2f} to ${k['rpc5_max']:.2f},
   so their revenue mostly follows how many customers they have. Across all {k['n_countries']} countries it ranges from
   ${k['rpc_min']:.2f} to ${k['rpc_max']:.2f}, but {k['one_customer_countries']} countries have a single customer, so the small-country figures are noisy.
4. **Genre revenue is concentrated.** {k['top_genre']} earns {k['top_genre_share']:.1f}% and the top three of {k['n_genres']} genres
   earn {k['top3_genre_share']:.1f}%.
5. **{never:.1f}% of the catalogue never sold** ({k['never_sold']:,} of {k['catalogue']:,} tracks). The data has no cost
   information, so what that inventory costs to carry cannot be assessed here.
6. **Most customers keep buying.** Of the {int(coh[coh['cohort_year'] == first]['customers_in_cohort'].iloc[0])} customers whose first purchase was in {first},
   {', '.join(f"{v:.0f}% were active in {a}" for a, v in zip(later['activity_year'], later['retention_pct']))}.
   ("Active" means at least one invoice that year, so a customer can skip a year and return.)
7. **RFM segments:** {', '.join(f'{n} {s.lower()}' for s, n in k['segments'].items())} customers. {k['freq_same']} of {k['customers']} customers have exactly
   {k['freq_mode']} invoices, so the Frequency score separates them almost not at all and the segments are really Recency plus Monetary;
   treat them as illustrative.

## Results by query

Top rows of the most useful outputs (all results are in `reports/results/` and the workbook).

**Revenue by year**

{md_table(r['yearly_revenue'])}

**Top 8 countries**

{md_table(r['revenue_by_country'].head(8))}

**Top 8 genres**

{md_table(r['revenue_by_genre'].head(8))}

## What this analysis cannot tell you

- The data is synthetic-style sample data, so patterns here do not describe any real store.
- There are no costs, margins, discounts or refunds, so this is revenue, not profit.
- Five years and {k['customers']} customers is too little for statistical claims; differences of a few percent are noise.
- RFM bands here are illustrative for the reason in point 7.
"""
    path.write_text(text)


def main() -> None:
    (REPORTS / "results").mkdir(parents=True, exist_ok=True)
    conn = connect()
    results = run_queries(conn)
    conn.close()
    for name, df in results.items():
        df.to_csv(REPORTS / "results" / f"{name}.csv", index=False)
        print(f"  {name:22s} {len(df):4d} rows")
    write_excel(results, REPORTS / "sales_analysis.xlsx")
    write_figures(results, REPORTS / "figures")
    write_findings(results, REPORTS / "findings.md")
    print("done")


if __name__ == "__main__":
    main()
